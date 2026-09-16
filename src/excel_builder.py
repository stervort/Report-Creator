"""
Excel Builder Module
Generates Excel files with transaction details
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class ExcelBuilder:
    """Build professional Excel files with transaction details"""
    
    def __init__(self, filepath):
        self.filepath = filepath
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Transactions"
        self.row_num = 1
    
    def add_header(self, title, company_name):
        """Add header section"""
        self.ws[f'A{self.row_num}'] = company_name
        self.ws[f'A{self.row_num}'].font = Font(size=14, bold=True, color="FFFFFF")
        self.ws[f'A{self.row_num}'].fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")
        self.ws.merge_cells(f'A{self.row_num}:E{self.row_num}')
        self.row_num += 1
        
        self.ws[f'A{self.row_num}'] = title
        self.ws[f'A{self.row_num}'].font = Font(size=12, bold=True)
        self.ws.merge_cells(f'A{self.row_num}:E{self.row_num}')
        self.row_num += 1
        
        self.ws[f'A{self.row_num}'] = f"Generated: {datetime.now().strftime('%m/%d/%Y %I:%M %p')}"
        self.ws[f'A{self.row_num}'].font = Font(size=9, italic=True)
        self.ws.merge_cells(f'A{self.row_num}:E{self.row_num}')
        self.row_num += 2
    
    def add_transaction_table(self, transactions, department=None):
        """Add transaction detail table"""
        if department:
            self.ws[f'A{self.row_num}'] = f"Department: {department}"
            self.ws[f'A{self.row_num}'].font = Font(bold=True, size=11)
            self.ws.merge_cells(f'A{self.row_num}:E{self.row_num}')
            self.row_num += 1
        
        headers = ['Date', 'Description', 'Amount', 'Name', 'Type']
        for col_num, header in enumerate(headers, 1):
            cell = self.ws.cell(row=self.row_num, column=col_num)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        self.row_num += 1
        
        for tx in transactions:
            self.ws.cell(row=self.row_num, column=1).value = tx.get('date')
            self.ws.cell(row=self.row_num, column=2).value = tx.get('description', '')
            self.ws.cell(row=self.row_num, column=3).value = tx.get('amount')
            self.ws.cell(row=self.row_num, column=4).value = tx.get('name', '')
            self.ws.cell(row=self.row_num, column=5).value = tx.get('type', '')
            
            self.ws.cell(row=self.row_num, column=3).number_format = '$#,##0.00'
            
            if tx.get('condensed', False):
                for col in range(1, 6):
                    self.ws.cell(row=self.row_num, column=col).font = Font(
                        italic=True, color="999999", size=9
                    )
            
            self.row_num += 1
        
        self.row_num += 1
    
    def set_column_widths(self):
        """Set optimal column widths"""
        self.ws.column_dimensions['A'].width = 15
        self.ws.column_dimensions['B'].width = 40
        self.ws.column_dimensions['C'].width = 15
        self.ws.column_dimensions['D'].width = 25
        self.ws.column_dimensions['E'].width = 12
    
    def save(self):
        """Save the workbook"""
        logger.info(f"Saving Excel file: {self.filepath}")
        self.wb.save(self.filepath)
        logger.info(f"Excel file created: {self.filepath}")


def generate_detail_excel(excel_path, report_name, company_name, transactions, departments=None):
    """Generate Excel file with transaction details"""
    builder = ExcelBuilder(excel_path)
    
    builder.add_header(report_name, company_name)
    
    if isinstance(transactions, dict):
        for dept, dept_txs in transactions.items():
            builder.add_transaction_table(dept_txs, department=dept)
    else:
        builder.add_transaction_table(transactions)
    
    builder.set_column_widths()
    builder.save()
