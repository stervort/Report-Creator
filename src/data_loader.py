"""
Data Loader Module
Loads and parses QB export files (Budget vs Actuals and General Ledger)
"""

import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class BudgetVsActualsLoader:
    """Load and parse Budget vs Actuals report from QB"""
    
    def __init__(self, filepath):
        self.filepath = filepath
        self.wb = openpyxl.load_workbook(filepath)
        self.ws = self.wb.active
        self.data = None
        self.classes = {}
        self.line_items = []
        
    def load(self):
        """Load and parse the BvA file"""
        logger.info(f"Loading Budget vs Actuals from {self.filepath}")
        
        # Extract class headers (row 5)
        self._extract_classes()
        
        # Extract line items and data
        self._extract_line_items()
        
        logger.info(f"Found {len(self.classes)} departments and {len(self.line_items)} line items")
        return self
    
    def _extract_classes(self):
        """Extract department/class names from row 5"""
        for col in range(2, 154):
            cell_val = self.ws.cell(row=5, column=col).value
            if cell_val and str(cell_val).strip() and 'Unnamed' not in str(cell_val):
                if cell_val not in self.classes:
                    self.classes[cell_val] = col
    
    def _extract_line_items(self):
        """Extract line items from column A"""
        for row in range(7, 230):
            cell_val = self.ws.cell(row=row, column=1).value
            if cell_val and not str(cell_val).strip().startswith('='):
                self.line_items.append({
                    'row': row,
                    'description': str(cell_val).strip()
                })
    
    def get_classes(self):
        """Return list of all department classes"""
        return list(self.classes.keys())
    
    def get_class_data(self, class_name):
        """Extract budget and actual data for a specific class"""
        if class_name not in self.classes:
            return None
        
        col = self.classes[class_name]
        
        data = {
            'class': class_name,
            'line_items': [],
        }
        
        for item in self.line_items:
            row = item['row']
            description = item['description']
            
            actual = self.ws.cell(row=row, column=col).value
            budget = self.ws.cell(row=row, column=col+1).value
            
            if isinstance(actual, str) and actual.startswith('='):
                actual = None
            if isinstance(budget, str) and budget.startswith('='):
                budget = None
            
            data['line_items'].append({
                'description': description,
                'actual': actual,
                'budget': budget,
            })
        
        return data
    
    def get_data_range(self, class_name, start_row, end_row):
        """Extract data range for a class"""
        if class_name not in self.classes:
            return None
        
        col = self.classes[class_name]
        
        data = {
            'class': class_name,
            'rows': []
        }
        
        for row in range(start_row, end_row + 1):
            description = self.ws.cell(row=row, column=1).value
            actual = self.ws.cell(row=row, column=col).value
            budget = self.ws.cell(row=row, column=col+1).value
            
            if description:
                data['rows'].append({
                    'row': row,
                    'description': str(description).strip() if description else '',
                    'actual': actual,
                    'budget': budget,
                })
        
        return data


class GeneralLedgerLoader:
    """Load and parse General Ledger from QB"""
    
    def __init__(self, filepath):
        self.filepath = filepath
        self.wb = openpyxl.load_workbook(filepath)
        self.ws = self.wb.active
        self.transactions = []
        self.classes = set()
        
    def load(self):
        """Load and parse the GL file"""
        logger.info(f"Loading General Ledger from {self.filepath}")
        
        for row in range(10, self.ws.max_row + 1):
            amount = self.ws.cell(row=row, column=9).value
            cls = self.ws.cell(row=row, column=11).value
            
            if amount and cls and 'Item class' not in str(cls):
                date = self.ws.cell(row=row, column=3).value
                tx_type = self.ws.cell(row=row, column=4).value
                name = self.ws.cell(row=row, column=6).value
                description = self.ws.cell(row=row, column=7).value
                
                self.transactions.append({
                    'row': row,
                    'date': date,
                    'type': tx_type,
                    'name': name,
                    'description': description,
                    'amount': amount,
                    'class': cls
                })
                
                self.classes.add(cls)
        
        logger.info(f"Loaded {len(self.transactions)} transactions from {len(self.classes)} classes")
        return self
    
    def get_transactions_by_class(self, class_name):
        """Get all transactions for a specific class"""
        return [t for t in self.transactions if t['class'] == class_name]
    
    def get_transactions_for_classes(self, class_list):
        """Get all transactions for multiple classes"""
        return [t for t in self.transactions if t['class'] in class_list]
    
    def get_classes(self):
        """Return list of all classes with transactions"""
        return sorted(list(self.classes))


def load_qb_data(bva_path, gl_path):
    """Load both QB files"""
    bva = BudgetVsActualsLoader(bva_path)
    bva.load()
    
    gl = GeneralLedgerLoader(gl_path)
    gl.load()
    
    return bva, gl
